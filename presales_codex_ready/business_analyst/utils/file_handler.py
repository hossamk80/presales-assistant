"""
utils/file_handler.py — File I/O: Extraction & Export
"""
import glob
import os
import streamlit as st
import pandas as pd
from io import BytesIO
from typing import List, Optional

from utils.document_blocks import parse_blocks
from utils import submission as _submission

# أقل عدد أحرف في الصفحة يُعتبر معه استخراج النص ناجحاً.
# ما دون ذلك يرجّح أن الصفحة صورة ممسوحة ضوئياً.
MIN_CHARS_PER_PAGE = 40


def ocr_available() -> bool:
    """هل أدوات الـ OCR الاختيارية مثبّتة؟"""
    try:
        import pytesseract  # noqa: F401
        from pdf2image import convert_from_bytes  # noqa: F401
    except ImportError:
        return False
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _ocr_pdf(file_bytes: bytes, file_name: str) -> str:
    """
    تشغيل OCR على ملف PDF ممسوح ضوئياً (عربي + إنجليزي).
    يتطلب: pytesseract + pdf2image + tesseract-ocr مع حزمة اللغة العربية.
    """
    import pytesseract
    from pdf2image import convert_from_bytes

    pages = convert_from_bytes(file_bytes, dpi=300)
    out = []
    progress = st.progress(0.0, text=f"OCR — {file_name}")
    for i, page in enumerate(pages, start=1):
        out.append(pytesseract.image_to_string(page, lang="ara+eng"))
        progress.progress(i / len(pages), text=f"OCR — {file_name} ({i}/{len(pages)})")
    progress.empty()
    return "\n".join(out)


def extract_texts_per_file(files: list) -> dict:
    """
    يستخرج نص كل ملف على حدة: {اسم الملف: النص}.

    الفصل ضروري لتصنيف المرفقات حسب دورها (كراسة / ملحق / كميات) بدل
    التعامل معها ككتلة واحدة.
    """
    return {f.name: _extract_single(f) for f in files}


def extract_text_from_files(files: list) -> tuple[str, list]:
    """
    Extract text from uploaded files. Returns (combined_text, list_of_filenames).
    Handles: PDF (with OCR fallback), DOCX/DOC, XLSX/XLS/CSV, HTML, TXT
    """
    per_file = extract_texts_per_file(files)
    combined = "\n".join(
        f"\n\n=== {name} ===\n{text}" for name, text in per_file.items()
    )
    return combined.strip(), list(per_file)


def _extract_single(file) -> str:
    """نص ملف واحد بلا ترويسة اسم الملف."""
    text_parts = []

    for file in [file]:
        ext = file.name.rsplit(".", 1)[-1].lower()
        try:
            if ext == "pdf":
                from pypdf import PdfReader

                raw = file.getvalue()
                reader = PdfReader(BytesIO(raw))
                page_count = len(reader.pages)
                pages_text = []
                for page in reader.pages:
                    t = page.extract_text()
                    if t:
                        pages_text.append(t)
                extracted = "\n".join(pages_text)

                # كشف الملفات الممسوحة ضوئياً: نص ضئيل مقارنة بعدد الصفحات
                is_scanned = page_count > 0 and len(extracted.strip()) < MIN_CHARS_PER_PAGE * page_count

                if is_scanned:
                    if ocr_available():
                        st.info(f"🔍 `{file.name}` يبدو ممسوحاً ضوئياً — جاري تشغيل OCR ({page_count} صفحة)...")
                        try:
                            extracted = _ocr_pdf(raw, file.name)
                            st.success(f"✅ اكتمل OCR لـ `{file.name}` — {len(extracted):,} حرف.")
                        except Exception as ocr_err:
                            st.error(f"❌ فشل OCR لـ `{file.name}`: {ocr_err}")
                    else:
                        st.error(
                            f"🚨 `{file.name}` ملف PDF ممسوح ضوئياً ولم يُستخرَج منه نص يُذكر "
                            f"({len(extracted.strip()):,} حرف من {page_count} صفحة). "
                            "**سيُحلَّل هذا الملف ناقصاً.** لتفعيل OCR ثبّت: "
                            "`pip install pytesseract pdf2image` + حزمة النظام "
                            "`tesseract-ocr tesseract-ocr-ara poppler-utils`."
                        )

                text_parts.append(extracted)

            elif ext in ("docx", "doc"):
                import docx2txt
                text_parts.append(docx2txt.process(file))

            elif ext in ("xlsx", "xls"):
                sheets = pd.read_excel(file, sheet_name=None)
                parts = [
                    f"--- ورقة: {name} ---\n{df.to_string(index=False)}"
                    for name, df in sheets.items()
                ]
                text_parts.append("\n\n".join(parts))

            elif ext == "csv":
                df = pd.read_csv(file)
                text_parts.append(df.to_string(index=False))

            elif ext in ("html", "htm"):
                from bs4 import BeautifulSoup
                content = file.getvalue().decode("utf-8", errors="replace")
                soup = BeautifulSoup(content, "html.parser")
                text_parts.append(soup.get_text(separator="\n"))

            elif ext == "txt":
                content = file.getvalue().decode("utf-8", errors="replace")
                text_parts.append(content)

            elif ext in ("png", "jpg", "jpeg", "gif", "webp"):
                text_parts.append("[صورة — لا يمكن استخراج النص منها تلقائياً]")

            else:
                text_parts.append(f"[تنسيق غير مدعوم: .{ext}]")

        except Exception as e:
            st.warning(f"⚠️ فشل قراءة `{file.name}`: {e}")
            text_parts.append(f"[فشل الاستخراج: {e}]")

    return "\n".join(text_parts).strip()


# ══════════════════════════════════════════════════════════════════════════════
#  بناء مستند Word
# ══════════════════════════════════════════════════════════════════════════════


# الهوية البصرية الافتراضية للمستندات المصدَّرة.
BRAND_COLOR = "003366"          # أزرق مؤسسي
BRAND_FONT_AR = "Traditional Arabic"
BRAND_FONT_EN = "Calibri"


# ─── نصوص المستند الثابتة ─────────────────────────────────────────────────────
#
# هذه نصوص تُكتب **داخل الملف المصدَّر**، فتتبع لغة المخرجات لا لغة الواجهة:
# مستخدم يعمل بواجهة عربية ويولّد عرضاً إنجليزياً يجب أن يخرج ملفه إنجليزياً
# بالكامل. لذلك لا تمر من `i18n` — تلك للواجهة وحدها.

_CONFIDENTIALITY = {
    "ar": (
        "هذا المستند سري للغاية ومُعدّ حصرياً للجهة المُرسَل إليها.\n"
        "الشركة المُقدِّمة: {company}\n"
        "يُحظر توزيع هذا المستند أو إعادة إنتاجه دون إذن كتابي مسبق."
    ),
    "en": (
        "This document is strictly confidential and prepared solely for the "
        "receiving entity.\n"
        "Submitted by: {company}\n"
        "Distribution or reproduction without prior written consent is prohibited."
    ),
}

# الرموز التي تُستبدل بقيم حقيقية قبل التصدير. تُكتب بين أقواس مربعة عمداً
# ليحجبها فحص النص النائب إن بقيت بلا قيمة — اسم شركة مفقود على الغلاف عيب
# لا يقلّ عن نص نائب منسي.
COMPANY_TOKENS = ("[اسم الشركة]", "[Company Name]", "[COMPANY]")


def confidentiality_notice(company_name: str, language: str = "ar") -> str:
    """
    إشعار السرية ومعلومات المستند بلغة المخرجات.

    كان مكتوباً عربياً في `views/doc_builder.py`، فكان العرض الإنجليزي الكامل
    يحمل صفحة عربية أمام لجنة الفتح.
    """
    company = str(company_name or "").strip()
    if language == "both":
        return (
            _CONFIDENTIALITY["ar"].format(company=company)
            + "\n\n"
            + _CONFIDENTIALITY["en"].format(company=company)
        )
    template = _CONFIDENTIALITY.get(language, _CONFIDENTIALITY["ar"])
    return template.format(company=company)


def resolve_document_tokens(text: str, company_name: str) -> str:
    """
    يستبدل رموز الشركة في نص المستند بقيمتها الحقيقية.

    خطاب التقديم الافتراضي يحمل `[اسم الشركة]`، وبوابة التصدير تحجب أي نص
    فيه `[...]`. فكان المستخدم الجديد يجد التصدير محجوباً بنص افتراضي شحنّاه
    نحن. الاستبدال هنا يرفع الحجب متى عُرف اسم الشركة، ويُبقيه متى جُهل —
    وهو الحجب الصحيح.
    """
    company = str(company_name or "").strip()
    if not company:
        return text
    out = str(text or "")
    for token in COMPANY_TOKENS:
        out = out.replace(token, company)
    return out


def _rgb(hex_color: str):
    """تحويل لون سداسي عشري إلى RGBColor مع تجاهل أي صيغة غير صالحة."""
    from docx.shared import RGBColor

    value = str(hex_color or "").lstrip("#").strip()
    try:
        return RGBColor.from_string(value.upper())
    except (ValueError, AttributeError):
        return RGBColor.from_string(BRAND_COLOR)


def _set_run_font(run, font_name: str, rtl: bool):
    """
    تثبيت الخط على مستوى XML.

    ضبط run.font.name وحده لا يكفي للعربية: Word يعامل العربية كنص معقّد
    (complex script) فيقرأ الخط من w:cs لا من w:ascii، فيظهر النص العربي بخط
    آخر رغم ضبط الاسم. نكتب السمات الأربع معاً.
    """
    from docx.oxml.ns import qn

    run.font.name = font_name
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        from docx.oxml import OxmlElement
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), font_name)
    if rtl:
        from docx.oxml import OxmlElement
        mark = OxmlElement("w:rtl")
        mark.set(qn("w:val"), "1")
        rPr.append(mark)


def _apply_brand_styles(doc, rtl: bool, brand_color: str, font_name: str):
    """
    تطبيق الهوية على أنماط المستند لا على كل فقرة.

    الأنماط تسري على الفهرس والقوائم والجداول التي يولّدها Word لاحقاً، وهي
    أشياء لا نمرّ عليها فقرةً فقرةً، فضبط الأنماط يغطيها كلها.
    """
    from docx.oxml.ns import qn
    from docx.shared import Pt

    color = _rgb(brand_color)

    def _style_font(style, size: Optional[int] = None, bold: Optional[bool] = None,
                    colored: bool = False):
        style.font.name = font_name
        if size is not None:
            style.font.size = Pt(size)
        if bold is not None:
            style.font.bold = bold
        if colored:
            style.font.color.rgb = color
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            from docx.oxml import OxmlElement
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(attr), font_name)

    try:
        _style_font(doc.styles["Normal"], size=14 if rtl else 11)
    except KeyError:
        pass

    for name, size in (("Heading 1", 20), ("Heading 2", 16), ("Heading 3", 14),
                       ("Title", 28)):
        try:
            _style_font(doc.styles[name], size=size, bold=True, colored=True)
        except KeyError:
            continue


def _set_rtl(paragraph):
    """ضبط اتجاه الفقرة من اليمين لليسار على مستوى XML."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    for run in paragraph.runs:
        rPr = run._r.get_or_add_rPr()
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        rPr.append(rtl)


def _add_field(paragraph, instruction: str):
    """إدراج حقل Word (مثل PAGE أو TOC) داخل فقرة."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, end):
        run._r.append(el)


def _enable_update_fields(doc):
    """يجعل Word يحدّث الفهرس تلقائياً عند فتح الملف."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)


def _add_page_numbers(doc):
    """ترقيم الصفحات في تذييل كل مقطع."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.text = ""
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_field(footer, " PAGE ")


def _para_dir(paragraph, rtl: bool, center: bool = False):
    """يضبط محاذاة الفقرة واتجاهها حسب لغة المخرجات."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT if rtl else WD_ALIGN_PARAGRAPH.LEFT
    if rtl:
        _set_rtl(paragraph)
    return paragraph


def _render_blocks_docx(doc, markdown: str, base_level: int = 1, rtl: bool = True):
    """كتابة كتل Markdown في مستند Word مع الحفاظ على البنية."""
    for block in parse_blocks(markdown):
        btype = block["type"]

        if btype == "heading":
            level = min(base_level + block["level"] - 1, 9)
            _para_dir(doc.add_heading(block["text"], level=level), rtl)

        elif btype == "paragraph":
            _para_dir(doc.add_paragraph(block["text"]), rtl)

        elif btype in ("bullets", "numbered"):
            style = "List Bullet" if btype == "bullets" else "List Number"
            for item in block["items"]:
                try:
                    p = doc.add_paragraph(item, style=style)
                except KeyError:
                    p = doc.add_paragraph(("\u2022 " if btype == "bullets" else "") + item)
                _para_dir(p, rtl)

        elif btype == "table":
            _add_docx_table(doc, block["header"], block["rows"], rtl=rtl)


def _add_docx_table(doc, header: list, rows: list, rtl: bool = True):
    """
    جدول Word بترويسة. في العربية تُقلب الأعمدة ليبدأ العمود الأول من اليمين.
    """
    if not header:
        return
    if rtl:
        header = list(reversed(header))
        rows = [list(reversed(r)) for r in rows]

    table = doc.add_table(rows=1, cols=len(header))
    try:
        table.style = "Light Shading Accent 1"
    except KeyError:
        table.style = "Table Grid"

    for i, col in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = str(col)
        for p in cell.paragraphs:
            _para_dir(p, rtl)
            for run in p.runs:
                run.bold = True

    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row[: len(header)]):
            cells[i].text = "" if val is None else str(val)
            for p in cells[i].paragraphs:
                _para_dir(p, rtl)
    doc.add_paragraph()


def _shade_cell(cell, hex_color: str):
    """تظليل خانة جدول Word — لا واجهة عليا لها في python-docx."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shading)


def _gantt_header(grid: dict, rtl: bool) -> list:
    """ترويسة المخطط: عمود الاسم ثم أرقام الوحدات الزمنية."""
    unit = grid.get("unit", "week")
    if rtl:
        label = "المرحلة"
        prefix = "أ" if unit == "week" else "ش"
    else:
        label = "Phase"
        prefix = "W" if unit == "week" else "M"
    return [label] + [f"{prefix}{i + 1}" for i in range(grid["columns"])]


def _add_gantt_docx(doc, grid: dict, rtl: bool, accent_hex: str):
    """
    يرسم المخطط الزمني جدولاً مظلَّلاً في Word.

    جدول لا صورة: يخرج نفسه في Word و PDF بلا مكتبة رسم ولا خط مفقود، ويبقى
    مقروءاً عند الطباعة بالأبيض والأسود.
    """
    header = _gantt_header(grid, rtl)
    # كل صف: نصّه وعلامة التظليل لكل خانة، محاذيان دائماً. القلب في العربية
    # يقلبهما معاً — قلب أحدهما وحده يُظلّل الأسبوع الخطأ.
    rows = [
        ([r["label"]] + ["" for _ in r["cells"]], [False] + list(r["cells"]))
        for r in grid["rows"]
    ]

    if rtl:
        header = list(reversed(header))
        rows = [(list(reversed(text)), list(reversed(marks))) for text, marks in rows]

    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"

    for i, col in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = str(col)
        for p in cell.paragraphs:
            _para_dir(p, rtl, center=True)
            for run in p.runs:
                run.bold = True

    fill = str(accent_hex or BRAND_COLOR).lstrip("#").upper()
    for text, marks in rows:
        cells = table.add_row().cells
        for i, (value, active) in enumerate(zip(text, marks)):
            cells[i].text = str(value)
            for p in cells[i].paragraphs:
                _para_dir(p, rtl, center=not value)
            if active:
                _shade_cell(cells[i], fill)
    doc.add_paragraph()


def _render_timeline_docx(doc, df_timeline, rtl: bool, brand_color: str,
                          empty_note: str):
    """جدول المراحل ثم المخطط المظلَّل — بلا الأعمدة المالية."""
    from utils.timeline import export_df, gantt_grid

    block = _df_to_table_block(export_df(df_timeline))
    if not block:
        _para_dir(doc.add_paragraph(empty_note), rtl)
        return
    _add_docx_table(doc, *block, rtl=rtl)
    grid = gantt_grid(df_timeline)
    if grid:
        _add_gantt_docx(doc, grid, rtl, brand_color)


def _df_to_table_block(df: Optional[pd.DataFrame]) -> Optional[tuple]:
    if df is None or df.empty:
        return None
    header = [str(c) for c in df.columns]
    rows = [["" if pd.isna(v) else str(v) for v in row] for row in df.itertuples(index=False)]
    return header, rows


def build_word_document(
    company_name: str,
    sections: list,
    template_bytes: Optional[bytes] = None,
    df_compliance: Optional[pd.DataFrame] = None,
    df_boq: Optional[pd.DataFrame] = None,
    df_timeline: Optional[pd.DataFrame] = None,
    include_toc: bool = True,
    include_page_numbers: bool = True,
    rtl: bool = True,
    proposal_title: str = "",
    entity_name: str = "",
    brand_color: str = BRAND_COLOR,
    font_name: str = "",
) -> BytesIO:
    """
    يبني مستند Word من قائمة أقسام مرتّبة.

    Args:
        sections: [{"key","title","kind","content"}] بالترتيب النهائي.
                  kind: cover · docinfo · ai · table_compliance · table_boq
        template_bytes: قالب Word للشركة يُحقن المحتوى بعده.
        proposal_title: عنوان العرض على الغلاف؛ يعود لعنوان عام إن كان فارغاً.
        entity_name: الجهة المصدِرة للمنافسة — تظهر كسطر "مقدَّم إلى".
        brand_color: لون العناوين سداسي عشري.
        font_name: خط المستند؛ يُختار حسب اللغة إن تُرك فارغاً.
    """
    from docx import Document
    from docx.shared import Cm, Pt

    font_name = font_name or (BRAND_FONT_AR if rtl else BRAND_FONT_EN)
    doc_title = proposal_title.strip() or ("العرض الفني" if rtl else "Technical Proposal")
    submitted_by = "مقدَّم من" if rtl else "Submitted by"
    submitted_to = "مقدَّم إلى" if rtl else "Submitted to"
    date_label = "التاريخ" if rtl else "Date"
    toc_title = "فهرس المحتويات" if rtl else "Table of Contents"
    toc_note = (
        "(إذا ظهر الفهرس فارغاً، اضغط داخله ثم F9 لتحديثه.)" if rtl
        else "(If the table of contents is empty, click inside it and press F9.)"
    )
    empty_note = "[هذا القسم فارغ]" if rtl else "[This section is empty]"

    if template_bytes:
        doc = Document(BytesIO(template_bytes))
        if len(doc.paragraphs) > 0:
            doc.add_page_break()
    else:
        doc = Document()
        for section in doc.sections:
            section.top_margin = Cm(2.5)
            section.bottom_margin = Cm(2.5)
            section.left_margin = Cm(3)
            section.right_margin = Cm(3)

    # قالب الشركة يحمل هويتها البصرية، فلا نفرض هويتنا فوقه.
    if not template_bytes:
        _apply_brand_styles(doc, rtl, brand_color, font_name)

    # ── صفحة الغلاف ───────────────────────────────────────────────────────────
    cover = _para_dir(doc.add_paragraph(), rtl, center=True)
    title_run = cover.add_run(doc_title)
    title_run.bold = True
    title_run.font.size = Pt(28)
    title_run.font.color.rgb = _rgb(brand_color)
    _set_run_font(title_run, font_name, rtl)

    for label, value in ((submitted_to, entity_name), (submitted_by, company_name)):
        if not str(value).strip():
            continue
        line = _para_dir(doc.add_paragraph(), rtl, center=True)
        run = line.add_run(f"{label}: {value}")
        run.font.size = Pt(16)
        _set_run_font(run, font_name, rtl)

    # ب-4: التقويمان معاً وكلٌّ موسوم — انظر `submission.cover_date`
    date_line = _para_dir(doc.add_paragraph(), rtl, center=True)
    date_run = date_line.add_run(f"{date_label}: {_submission.cover_date(rtl=rtl)}")
    date_run.font.size = Pt(12)
    _set_run_font(date_run, font_name, rtl)
    doc.add_page_break()

    # ── فهرس المحتويات ────────────────────────────────────────────────────────
    if include_toc:
        _para_dir(doc.add_heading(toc_title, 1), rtl)
        _add_field(doc.add_paragraph(), r' TOC \o "1-3" \h \z \u ')
        _para_dir(doc.add_paragraph(toc_note), rtl)
        _enable_update_fields(doc)
        doc.add_page_break()

    # ── الأقسام ───────────────────────────────────────────────────────────────
    for idx, sec in enumerate(sections):
        kind = sec.get("kind")
        _para_dir(doc.add_heading(sec.get("title", ""), 1), rtl)

        if kind == "table_timeline":
            _render_timeline_docx(doc, df_timeline, rtl, brand_color, empty_note)
        elif kind in ("table_compliance", "table_boq"):
            block = _df_to_table_block(
                df_compliance if kind == "table_compliance" else df_boq
            )
            if block:
                _add_docx_table(doc, *block, rtl=rtl)
            else:
                _para_dir(doc.add_paragraph(empty_note), rtl)
        else:
            content = (sec.get("content") or "").strip()
            if content:
                _render_blocks_docx(doc, content, base_level=2, rtl=rtl)
            else:
                _para_dir(doc.add_paragraph(empty_note), rtl)

        if idx < len(sections) - 1:
            doc.add_page_break()

    if include_page_numbers:
        _add_page_numbers(doc)

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


# ══════════════════════════════════════════════════════════════════════════════
#  بناء مستند PDF
# ══════════════════════════════════════════════════════════════════════════════

# خطوط مرشّحة تدعم الحروف العربية، بترتيب الأفضلية، عبر أنظمة التشغيل الشائعة.
_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/**/Amiri*.ttf",
    "/usr/share/fonts/**/NotoNaskhArabic*.ttf",
    "/usr/share/fonts/**/NotoSansArabic*.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/**/DejaVuSans.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "C:/Windows/Fonts/tahoma.ttf",
]


def _find_arabic_font() -> tuple[Optional[str], Optional[str]]:
    """يبحث عن خط يدعم العربية. يُرجع (المسار العادي، المسار العريض)."""
    for pattern in _FONT_CANDIDATES:
        matches = sorted(glob.glob(pattern, recursive=True))
        if not matches:
            continue
        regular = matches[0]
        stem, ext = os.path.splitext(regular)
        for suffix in ("-Bold", "Bold", "bd", "-bold"):
            candidate = f"{stem}{suffix}{ext}"
            if os.path.exists(candidate):
                return regular, candidate
        return regular, None
    return None, None


def _shape(text: str, rtl: bool = True) -> str:
    """تشكيل الحروف العربية وترتيبها بصرياً للعرض في PDF."""
    if not rtl:
        return str(text)

    import arabic_reshaper
    from bidi.algorithm import get_display

    return get_display(arabic_reshaper.reshape(str(text)))


def _wrap_shaped(text: str, font: str, size: float, max_width: float,
                 rtl: bool = True) -> str:
    """
    يلفّ النص يدوياً ثم يشكّل كل سطر على حدة.

    التشكيل قبل اللف يعطي ترتيباً بصرياً خاطئاً عند تقسيم reportlab للفقرة،
    لأن أول سطر بصري سيحمل آخر النص المنطقي. لذا نلفّ أولاً ثم نشكّل.
    """
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from xml.sax.saxutils import escape

    lines, current = [], ""
    for word in str(text).split():
        trial = f"{current} {word}".strip()
        if stringWidth(_shape(trial, rtl), font, size) <= max_width or not current:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return "<br/>".join(escape(_shape(ln, rtl)) for ln in lines) or "&nbsp;"


def _rtl_toc_class():
    """
    فهرس محتويات باتجاه من اليمين لليسار.

    فهرس reportlab القياسي يرسم رقم الصفحة عند الحافة اليمنى — وهو الموضع
    الصحيح في المستندات الإنجليزية، لكنه في العربية يقع فوق بداية العنوان.
    نستبدل بناء الجدول فقط: عمود ضيّق لرقم الصفحة على اليسار، وعمود العنوان
    على اليمين. باقي آلية إعادة البناء ترثها كما هي.
    """
    from reportlab.lib.enums import TA_LEFT, TA_RIGHT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, Spacer, Table
    from reportlab.platypus.tableofcontents import TableOfContents

    class _RtlToc(TableOfContents):
        def __init__(self, rtl: bool = True):
            super().__init__()
            self._rtl = rtl

        def wrap(self, availWidth, availHeight):
            entries = self._lastEntries or [(0, "", 0, None)]
            num_w = 1.6 * cm
            # في العربية رقم الصفحة على اليسار والعنوان على اليمين، والعكس بالإنجليزية
            widths = (num_w, availWidth - num_w) if self._rtl else (availWidth - num_w, num_w)
            rows = []

            for level, text, page_num, _key in entries:
                style = self.getLevelStyle(level)
                num_style = ParagraphStyle(
                    f"tocnum{level}", parent=style,
                    alignment=TA_LEFT if self._rtl else TA_RIGHT,
                    rightIndent=0, leftIndent=0,
                )
                if style.spaceBefore:
                    rows.append([Spacer(1, style.spaceBefore), Spacer(1, style.spaceBefore)])
                cells = [Paragraph(str(page_num), num_style), Paragraph(text, style)]
                if not self._rtl:
                    cells.reverse()
                rows.append(cells)

            self._table = Table(rows, colWidths=widths, style=self.tableStyle)
            self.width, self.height = self._table.wrapOn(self.canv, availWidth, availHeight)
            return self.width, self.height

    return _RtlToc


def build_pdf_document(
    company_name: str,
    sections: list,
    df_compliance: Optional[pd.DataFrame] = None,
    df_boq: Optional[pd.DataFrame] = None,
    include_toc: bool = True,
    include_page_numbers: bool = True,
    rtl: bool = True,
    proposal_title: str = "",
    entity_name: str = "",
    brand_color: str = BRAND_COLOR,
    df_timeline: Optional[pd.DataFrame] = None,
) -> BytesIO:
    """
    يبني نسخة PDF من نفس الأقسام. في العربية يُشكَّل النص ويُحاذى لليمين.

    الغلاف والعناوين تتبع نفس هوية Word حتى لا تختلف الصيغتان أمام لجنة الفتح.

    Raises:
        ImportError: إذا لم تكن مكتبات الـ PDF أو خط عربي متوفراً.
    """
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        BaseDocTemplate,
        Frame,
        PageBreak,
        PageTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )

    regular, bold = _find_arabic_font()
    if not regular:
        raise ImportError(
            "لم يُعثر على خط يدعم الحروف العربية على هذا النظام. "
            "ثبّت مثلاً: sudo apt install fonts-dejavu fonts-hosny-amiri"
        )

    FONT, FONT_B = "ArabicBody", "ArabicBold"
    pdfmetrics.registerFont(TTFont(FONT, regular))
    pdfmetrics.registerFont(TTFont(FONT_B, bold or regular))

    side = TA_RIGHT if rtl else TA_LEFT
    accent = colors.HexColor(f"#{str(brand_color or BRAND_COLOR).lstrip('#')}")

    styles = getSampleStyleSheet()
    body = ParagraphStyle("ArBody", parent=styles["Normal"], fontName=FONT,
                          fontSize=11, leading=19, alignment=side, spaceAfter=6)
    h1 = ParagraphStyle("ArH1", parent=body, fontName=FONT_B, fontSize=17,
                        leading=26, spaceBefore=10, spaceAfter=10, textColor=accent)
    h2 = ParagraphStyle("ArH2", parent=body, fontName=FONT_B, fontSize=14,
                        leading=22, spaceBefore=8, spaceAfter=6, textColor=accent)
    h3 = ParagraphStyle("ArH3", parent=body, fontName=FONT_B, fontSize=12,
                        leading=20, spaceBefore=6, spaceAfter=4, textColor=accent)
    title_style = ParagraphStyle("ArTitle", parent=body, fontName=FONT_B,
                                 fontSize=26, leading=36, alignment=TA_CENTER,
                                 textColor=accent)
    center = ParagraphStyle("ArCenter", parent=body, alignment=TA_CENTER)

    class _Doc(BaseDocTemplate):
        """يسجّل العناوين في الفهرس بعد رسم كل عنصر."""

        def afterFlowable(self, flowable):
            level = getattr(flowable, "_toc_level", None)
            if level is None:
                return
            self.notify("TOCEntry", (level, getattr(flowable, "_toc_text", ""), self.page))

    bio = BytesIO()
    margin = 2.2 * cm
    doc = _Doc(
        bio, pagesize=A4,
        leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin,
        title=" — ".join(filter(None, [
            proposal_title.strip() or ("العرض الفني" if rtl else "Technical Proposal"),
            company_name,
        ])),
    )
    avail = doc.width

    def on_page(canvas, doc_):
        if not include_page_numbers:
            return
        canvas.saveState()
        canvas.setFont(FONT, 9)
        canvas.drawCentredString(A4[0] / 2, margin / 2, str(doc_.page))
        canvas.restoreState()

    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="body")
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=on_page)])

    def P(text, style, width=None):
        return Paragraph(_wrap_shaped(text, style.fontName, style.fontSize,
                                      width or avail, rtl), style)

    def H(text, style, level):
        """عنوان يُسجَّل في فهرس المحتويات."""
        p = P(text, style)
        p._toc_level = level
        p._toc_text = _shape(text, rtl)
        return p

    story = []

    # ── صفحة الغلاف ───────────────────────────────────────────────────────────
    story.append(Spacer(1, 6 * cm))
    story.append(P(
        proposal_title.strip() or ("العرض الفني" if rtl else "Technical Proposal"),
        title_style,
    ))
    story.append(Spacer(1, 1 * cm))
    if entity_name:
        story.append(P(
            f"{'مقدَّم إلى' if rtl else 'Submitted to'}: {entity_name}", center))
    if company_name:
        story.append(P(
            f"{'مقدَّم من' if rtl else 'Submitted by'}: {company_name}", center))
    # ب-4: نفس سطر التاريخ في Word — الصيغتان لا تختلفان أمام لجنة الفتح
    story.append(P(
        f"{'التاريخ' if rtl else 'Date'}: "
        f"{_submission.cover_date(rtl=rtl)}", center))
    story.append(PageBreak())

    # ── الفهرس ────────────────────────────────────────────────────────────────
    if include_toc:
        story.append(P("فهرس المحتويات" if rtl else "Table of Contents", h1))
        toc = _rtl_toc_class()(rtl=rtl)
        toc.levelStyles = [
            ParagraphStyle("toc1", parent=body, fontName=FONT_B, fontSize=12,
                           leading=22, alignment=side, spaceBefore=4),
            ParagraphStyle("toc2", parent=body, fontName=FONT, fontSize=10.5,
                           leading=18, alignment=side,
                           **({"rightIndent": 20} if rtl else {"leftIndent": 20})),
        ]
        story.append(toc)
        story.append(PageBreak())

    def table_flowable(header: list, rows: list):
        if rtl:
            header = list(reversed(header))
            rows = [list(reversed(r)) for r in rows]
        col_w = avail / max(len(header), 1)
        cell = ParagraphStyle("cell", parent=body, fontSize=9, leading=14, spaceAfter=0)
        cell_b = ParagraphStyle("cellb", parent=cell, fontName=FONT_B)
        data = [[P(str(c), cell_b, col_w - 10) for c in header]]
        for row in rows:
            data.append([P("" if v is None else str(v), cell, col_w - 10)
                         for v in row[: len(header)]])
        t = Table(data, colWidths=[col_w] * len(header), repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), accent),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F1F5F9")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        return t

    def gantt_flowable(grid: dict):
        """
        المخطط الزمني جدولاً مظلَّلاً — نفس ما يخرج في Word.

        جدول لا صورة: بلا مكتبة رسم ولا خط مفقود، ويبقى مقروءاً عند الطباعة
        بالأبيض والأسود لأن الخانة المظلَّلة تظهر رمادية داكنة.
        """
        header = _gantt_header(grid, rtl)
        rows = [
            ([r["label"]] + ["" for _ in r["cells"]], [False] + list(r["cells"]))
            for r in grid["rows"]
        ]
        if rtl:
            header = list(reversed(header))
            rows = [(list(reversed(txt)), list(reversed(mk))) for txt, mk in rows]

        # عمود الاسم أعرض من خانات الزمن — الأخيرة فارغة يكفيها التظليل
        name_col = min(6 * cm, avail * 0.35)
        unit_w = (avail - name_col) / max(grid["columns"], 1)
        widths = [unit_w] * len(header)
        widths[len(header) - 1 if rtl else 0] = name_col

        cell = ParagraphStyle("gcell", parent=body, fontSize=8, leading=11,
                              spaceAfter=0, alignment=TA_CENTER)
        cell_b = ParagraphStyle("gcellb", parent=cell, fontName=FONT_B)

        data = [[P(str(c), cell_b, w - 6) for c, w in zip(header, widths)]]
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), accent),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 2),
            ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ]
        for r, (text, marks) in enumerate(rows, start=1):
            data.append([P(str(v), cell, w - 6) for v, w in zip(text, widths)])
            for c, active in enumerate(marks):
                if active:
                    style.append(("BACKGROUND", (c, r), (c, r), accent))

        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle(style))
        return t

    heading_styles = {1: h1, 2: h2, 3: h3}

    for idx, sec in enumerate(sections):
        kind = sec.get("kind")
        story.append(H(sec.get("title", ""), h1, level=0))

        if kind == "table_timeline":
            from utils.timeline import export_df, gantt_grid

            block = _df_to_table_block(export_df(df_timeline))
            if block:
                story.append(table_flowable(*block))
                grid = gantt_grid(df_timeline)
                if grid:
                    story.append(Spacer(1, 0.4 * cm))
                    story.append(gantt_flowable(grid))
            else:
                story.append(P(
                    "[لا توجد خطة زمنية]" if rtl else "[No timeline]", body))
        elif kind == "table_compliance":
            block = _df_to_table_block(df_compliance)
            story.append(table_flowable(*block) if block
                         else P("[لا توجد بيانات في جدول الامتثال]", body))
        elif kind == "table_boq":
            block = _df_to_table_block(df_boq)
            story.append(table_flowable(*block) if block
                         else P("[لا توجد بيانات في جدول الكميات]", body))
        else:
            content = (sec.get("content") or "").strip()
            if not content:
                story.append(P(
                    "[هذا القسم فارغ]" if rtl else "[This section is empty]", body))
            for block in parse_blocks(content):
                btype = block["type"]
                if btype == "heading":
                    style = heading_styles.get(min(block["level"], 3), h3)
                    story.append(H(block["text"], style, level=1))
                elif btype == "paragraph":
                    story.append(P(block["text"], body))
                elif btype in ("bullets", "numbered"):
                    for n, item in enumerate(block["items"], start=1):
                        marker = "•" if btype == "bullets" else f"{n}."
                        story.append(P(
                            f"{item} {marker}" if rtl else f"{marker} {item}", body))
                elif btype == "table":
                    story.append(table_flowable(block["header"], block["rows"]))
                    story.append(Spacer(1, 0.3 * cm))

        if idx < len(sections) - 1:
            story.append(PageBreak())

    # multiBuild يُشغّل بناءين حتى تُحلّ أرقام صفحات الفهرس
    doc.multiBuild(story)
    bio.seek(0)
    return bio

