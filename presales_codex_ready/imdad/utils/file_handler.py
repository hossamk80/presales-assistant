"""
utils/file_handler.py — File I/O: Extraction & Export
"""
import datetime
import glob
import os
import streamlit as st
import pandas as pd
from io import BytesIO
from typing import List, Optional

from utils.document_blocks import parse_blocks

# أقل عدد أحرف في الصفحة يُعتبر معه استخراج النص ناجحاً.
# ما دون ذلك يرجّح أن الصفحة صورة ممسوحة ضوئياً.
MIN_CHARS_PER_PAGE = 40


def ocr_available() -> bool:
    """هل أدوات الـ OCR الاختيارية مثبّتة؟"""
    try:
        import pytesseract  # noqa: F401
        from pdf2image import convert_from_bytes  # noqa: F401
    except ImportError:
        return False
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


def _ocr_pdf(file_bytes: bytes, file_name: str) -> str:
    """
    تشغيل OCR على ملف PDF ممسوح ضوئياً (عربي + إنجليزي).
    يتطلب: pytesseract + pdf2image + tesseract-ocr مع حزمة اللغة العربية.
    """
    import pytesseract
    from pdf2image import convert_from_bytes

    pages = convert_from_bytes(file_bytes, dpi=300)
    out = []
    progress = st.progress(0.0, text=f"OCR — {file_name}")
    for i, page in enumerate(pages, start=1):
        out.append(pytesseract.image_to_string(page, lang="ara+eng"))
        progress.progress(i / len(pages), text=f"OCR — {file_name} ({i}/{len(pages)})")
    progress.empty()
    return "\n".join(out)


def extract_text_from_files(files: list) -> tuple[str, list]:
    """
    Extract text from uploaded files. Returns (combined_text, list_of_filenames).
    Handles: PDF (with OCR fallback), DOCX/DOC, XLSX/XLS/CSV, HTML, TXT
    """
    text_parts = []
    file_names = []

    for file in files:
        file_names.append(file.name)
        ext = file.name.rsplit(".", 1)[-1].lower()
        try:
            if ext == "pdf":
                from pypdf import PdfReader

                raw = file.getvalue()
                reader = PdfReader(BytesIO(raw))
                page_count = len(reader.pages)
                pages_text = []
                for page in reader.pages:
                    t = page.extract_text()
                    if t:
                        pages_text.append(t)
                extracted = "\n".join(pages_text)

                # كشف الملفات الممسوحة ضوئياً: نص ضئيل مقارنة بعدد الصفحات
                is_scanned = page_count > 0 and len(extracted.strip()) < MIN_CHARS_PER_PAGE * page_count

                if is_scanned:
                    if ocr_available():
                        st.info(f"🔍 `{file.name}` يبدو ممسوحاً ضوئياً — جاري تشغيل OCR ({page_count} صفحة)...")
                        try:
                            extracted = _ocr_pdf(raw, file.name)
                            st.success(f"✅ اكتمل OCR لـ `{file.name}` — {len(extracted):,} حرف.")
                        except Exception as ocr_err:
                            st.error(f"❌ فشل OCR لـ `{file.name}`: {ocr_err}")
                    else:
                        st.error(
                            f"🚨 `{file.name}` ملف PDF ممسوح ضوئياً ولم يُستخرَج منه نص يُذكر "
                            f"({len(extracted.strip()):,} حرف من {page_count} صفحة). "
                            "**سيُحلَّل هذا الملف ناقصاً.** لتفعيل OCR ثبّت: "
                            "`pip install pytesseract pdf2image` + حزمة النظام "
                            "`tesseract-ocr tesseract-ocr-ara poppler-utils`."
                        )

                text_parts.append(f"\n\n=== {file.name} ===\n" + extracted)

            elif ext in ("docx", "doc"):
                import docx2txt
                text_parts.append(f"\n\n=== {file.name} ===\n" + docx2txt.process(file))

            elif ext in ("xlsx", "xls"):
                sheets = pd.read_excel(file, sheet_name=None)
                parts = [
                    f"--- ورقة: {name} ---\n{df.to_string(index=False)}"
                    for name, df in sheets.items()
                ]
                text_parts.append(f"\n\n=== {file.name} (Excel) ===\n" + "\n\n".join(parts))

            elif ext == "csv":
                df = pd.read_csv(file)
                text_parts.append(f"\n\n=== {file.name} (CSV) ===\n" + df.to_string(index=False))

            elif ext in ("html", "htm"):
                from bs4 import BeautifulSoup
                content = file.getvalue().decode("utf-8", errors="replace")
                soup = BeautifulSoup(content, "html.parser")
                text_parts.append(f"\n\n=== {file.name} (HTML) ===\n" + soup.get_text(separator="\n"))

            elif ext == "txt":
                content = file.getvalue().decode("utf-8", errors="replace")
                text_parts.append(f"\n\n=== {file.name} ===\n" + content)

            elif ext in ("png", "jpg", "jpeg", "gif", "webp"):
                text_parts.append(f"\n\n=== {file.name} ===\n[صورة — لا يمكن استخراج النص منها تلقائياً]")

            else:
                text_parts.append(f"\n\n=== {file.name} ===\n[تنسيق غير مدعوم: .{ext}]")

        except Exception as e:
            st.warning(f"⚠️ فشل قراءة `{file.name}`: {e}")
            text_parts.append(f"\n\n=== {file.name} ===\n[فشل الاستخراج: {e}]")

    return "\n".join(text_parts).strip(), file_names


# ══════════════════════════════════════════════════════════════════════════════
#  بناء مستند Word
# ══════════════════════════════════════════════════════════════════════════════


def _set_rtl(paragraph):
    """ضبط اتجاه الفقرة من اليمين لليسار على مستوى XML."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    for run in paragraph.runs:
        rPr = run._r.get_or_add_rPr()
        rtl = OxmlElement("w:rtl")
        rtl.set(qn("w:val"), "1")
        rPr.append(rtl)


def _add_field(paragraph, instruction: str):
    """إدراج حقل Word (مثل PAGE أو TOC) داخل فقرة."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for el in (begin, instr, separate, end):
        run._r.append(el)


def _enable_update_fields(doc):
    """يجعل Word يحدّث الفهرس تلقائياً عند فتح الملف."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        el = OxmlElement("w:updateFields")
        el.set(qn("w:val"), "true")
        settings.append(el)


def _add_page_numbers(doc):
    """ترقيم الصفحات في تذييل كل مقطع."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.text = ""
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_field(footer, " PAGE ")


def _render_blocks_docx(doc, markdown: str, base_level: int = 1):
    """كتابة كتل Markdown في مستند Word مع الحفاظ على البنية."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    for block in parse_blocks(markdown):
        btype = block["type"]

        if btype == "heading":
            level = min(base_level + block["level"] - 1, 9)
            p = doc.add_heading(block["text"], level=level)
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_rtl(p)

        elif btype == "paragraph":
            p = doc.add_paragraph(block["text"])
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_rtl(p)

        elif btype in ("bullets", "numbered"):
            style = "List Bullet" if btype == "bullets" else "List Number"
            for item in block["items"]:
                try:
                    p = doc.add_paragraph(item, style=style)
                except KeyError:
                    p = doc.add_paragraph(("• " if btype == "bullets" else "") + item)
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                _set_rtl(p)

        elif btype == "table":
            _add_docx_table(doc, block["header"], block["rows"])


def _add_docx_table(doc, header: list, rows: list):
    """جدول Word بترويسة، بترتيب أعمدة مقلوب ليناسب القراءة من اليمين."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if not header:
        return
    header = list(reversed(header))
    rows = [list(reversed(r)) for r in rows]

    table = doc.add_table(rows=1, cols=len(header))
    try:
        table.style = "Light Shading Accent 1"
    except KeyError:
        table.style = "Table Grid"

    for i, col in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = str(col)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _set_rtl(p)
            for run in p.runs:
                run.bold = True

    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row[: len(header)]):
            cells[i].text = "" if val is None else str(val)
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                _set_rtl(p)
    doc.add_paragraph()


def _df_to_table_block(df: Optional[pd.DataFrame]) -> Optional[tuple]:
    if df is None or df.empty:
        return None
    header = [str(c) for c in df.columns]
    rows = [["" if pd.isna(v) else str(v) for v in row] for row in df.itertuples(index=False)]
    return header, rows


def build_word_document(
    company_name: str,
    sections: list,
    template_bytes: Optional[bytes] = None,
    df_compliance: Optional[pd.DataFrame] = None,
    df_boq: Optional[pd.DataFrame] = None,
    include_toc: bool = True,
    include_page_numbers: bool = True,
) -> BytesIO:
    """
    يبني مستند Word من قائمة أقسام مرتّبة.

    Args:
        sections: [{"key","title","kind","content"}] بالترتيب النهائي.
                  kind: cover · docinfo · ai · table_compliance · table_boq
        template_bytes: قالب Word للشركة يُحقن المحتوى بعده.
    """
    from docx import Document
    from docx.shared import Cm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_SECTION

    if template_bytes:
        doc = Document(BytesIO(template_bytes))
        if len(doc.paragraphs) > 0:
            doc.add_page_break()
    else:
        doc = Document()
        for section in doc.sections:
            section.top_margin = Cm(2.5)
            section.bottom_margin = Cm(2.5)
            section.left_margin = Cm(3)
            section.right_margin = Cm(3)

    # ── صفحة العنوان ──────────────────────────────────────────────────────────
    title = doc.add_heading("العرض الفني", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_rtl(title)
    if company_name:
        sub = doc.add_paragraph(f"مقدَّم من: {company_name}")
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _set_rtl(sub)
    date_p = doc.add_paragraph(f"التاريخ: {datetime.date.today().strftime('%Y/%m/%d')}")
    date_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set_rtl(date_p)
    doc.add_page_break()

    # ── فهرس المحتويات ────────────────────────────────────────────────────────
    if include_toc:
        toc_head = doc.add_heading("فهرس المحتويات", 1)
        toc_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_rtl(toc_head)
        toc_p = doc.add_paragraph()
        _add_field(toc_p, r' TOC \o "1-3" \h \z \u ')
        note = doc.add_paragraph(
            "(إذا ظهر الفهرس فارغاً، اضغط داخله ثم F9 لتحديثه.)"
        )
        note.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_rtl(note)
        _enable_update_fields(doc)
        doc.add_page_break()

    # ── الأقسام ───────────────────────────────────────────────────────────────
    for idx, sec in enumerate(sections):
        kind = sec.get("kind")
        heading = doc.add_heading(sec.get("title", ""), 1)
        heading.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _set_rtl(heading)

        if kind == "table_compliance":
            block = _df_to_table_block(df_compliance)
            if block:
                _add_docx_table(doc, *block)
            else:
                p = doc.add_paragraph("[لا توجد بيانات في جدول الامتثال]")
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                _set_rtl(p)
        elif kind == "table_boq":
            block = _df_to_table_block(df_boq)
            if block:
                _add_docx_table(doc, *block)
            else:
                p = doc.add_paragraph("[لا توجد بيانات في جدول الكميات]")
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                _set_rtl(p)
        else:
            content = (sec.get("content") or "").strip()
            if content:
                _render_blocks_docx(doc, content, base_level=2)
            else:
                p = doc.add_paragraph("[هذا القسم فارغ]")
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                _set_rtl(p)

        if idx < len(sections) - 1:
            doc.add_page_break()

    if include_page_numbers:
        _add_page_numbers(doc)

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio


# ══════════════════════════════════════════════════════════════════════════════
#  بناء مستند PDF
# ══════════════════════════════════════════════════════════════════════════════

# خطوط مرشّحة تدعم الحروف العربية، بترتيب الأفضلية، عبر أنظمة التشغيل الشائعة.
_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/**/Amiri*.ttf",
    "/usr/share/fonts/**/NotoNaskhArabic*.ttf",
    "/usr/share/fonts/**/NotoSansArabic*.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/**/DejaVuSans.ttf",
    "/Library/Fonts/Arial Unicode.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "C:/Windows/Fonts/arial.ttf",
    "C:/Windows/Fonts/tahoma.ttf",
]


def _find_arabic_font() -> tuple[Optional[str], Optional[str]]:
    """يبحث عن خط يدعم العربية. يُرجع (المسار العادي، المسار العريض)."""
    for pattern in _FONT_CANDIDATES:
        matches = sorted(glob.glob(pattern, recursive=True))
        if not matches:
            continue
        regular = matches[0]
        stem, ext = os.path.splitext(regular)
        for suffix in ("-Bold", "Bold", "bd", "-bold"):
            candidate = f"{stem}{suffix}{ext}"
            if os.path.exists(candidate):
                return regular, candidate
        return regular, None
    return None, None


def _shape(text: str) -> str:
    """تشكيل الحروف العربية وترتيبها بصرياً للعرض في PDF."""
    import arabic_reshaper
    from bidi.algorithm import get_display

    return get_display(arabic_reshaper.reshape(str(text)))


def _wrap_shaped(text: str, font: str, size: float, max_width: float) -> str:
    """
    يلفّ النص يدوياً ثم يشكّل كل سطر على حدة.

    التشكيل قبل اللف يعطي ترتيباً بصرياً خاطئاً عند تقسيم reportlab للفقرة،
    لأن أول سطر بصري سيحمل آخر النص المنطقي. لذا نلفّ أولاً ثم نشكّل.
    """
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from xml.sax.saxutils import escape

    lines, current = [], ""
    for word in str(text).split():
        trial = f"{current} {word}".strip()
        if stringWidth(_shape(trial), font, size) <= max_width or not current:
            current = trial
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return "<br/>".join(escape(_shape(ln)) for ln in lines) or "&nbsp;"


def _rtl_toc_class():
    """
    فهرس محتويات باتجاه من اليمين لليسار.

    فهرس reportlab القياسي يرسم رقم الصفحة عند الحافة اليمنى — وهو الموضع
    الصحيح في المستندات الإنجليزية، لكنه في العربية يقع فوق بداية العنوان.
    نستبدل بناء الجدول فقط: عمود ضيّق لرقم الصفحة على اليسار، وعمود العنوان
    على اليمين. باقي آلية إعادة البناء ترثها كما هي.
    """
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, Spacer, Table
    from reportlab.platypus.tableofcontents import TableOfContents

    class _RtlToc(TableOfContents):
        def wrap(self, availWidth, availHeight):
            entries = self._lastEntries or [(0, "", 0, None)]
            num_w = 1.6 * cm
            rows, styles = [], []

            for level, text, page_num, _key in entries:
                style = self.getLevelStyle(level)
                num_style = ParagraphStyle(
                    f"tocnum{level}", parent=style,
                    alignment=TA_LEFT, rightIndent=0, leftIndent=0,
                )
                if style.spaceBefore:
                    rows.append([Spacer(1, style.spaceBefore), Spacer(1, style.spaceBefore)])
                rows.append([
                    Paragraph(str(page_num), num_style),
                    Paragraph(text, style),
                ])

            self._table = Table(
                rows, colWidths=(num_w, availWidth - num_w), style=self.tableStyle
            )
            self.width, self.height = self._table.wrapOn(self.canv, availWidth, availHeight)
            return self.width, self.height

    return _RtlToc


def build_pdf_document(
    company_name: str,
    sections: list,
    df_compliance: Optional[pd.DataFrame] = None,
    df_boq: Optional[pd.DataFrame] = None,
    include_toc: bool = True,
    include_page_numbers: bool = True,
) -> BytesIO:
    """
    يبني نسخة PDF من نفس الأقسام، بتشكيل عربي صحيح ومحاذاة لليمين.

    Raises:
        ImportError: إذا لم تكن مكتبات الـ PDF أو خط عربي متوفراً.
    """
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        BaseDocTemplate,
        Frame,
        PageBreak,
        PageTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )

    regular, bold = _find_arabic_font()
    if not regular:
        raise ImportError(
            "لم يُعثر على خط يدعم الحروف العربية على هذا النظام. "
            "ثبّت مثلاً: sudo apt install fonts-dejavu fonts-hosny-amiri"
        )

    FONT, FONT_B = "ArabicBody", "ArabicBold"
    pdfmetrics.registerFont(TTFont(FONT, regular))
    pdfmetrics.registerFont(TTFont(FONT_B, bold or regular))

    styles = getSampleStyleSheet()
    body = ParagraphStyle("ArBody", parent=styles["Normal"], fontName=FONT,
                          fontSize=11, leading=19, alignment=TA_RIGHT, spaceAfter=6)
    h1 = ParagraphStyle("ArH1", parent=body, fontName=FONT_B, fontSize=17,
                        leading=26, spaceBefore=10, spaceAfter=10)
    h2 = ParagraphStyle("ArH2", parent=body, fontName=FONT_B, fontSize=14,
                        leading=22, spaceBefore=8, spaceAfter=6)
    h3 = ParagraphStyle("ArH3", parent=body, fontName=FONT_B, fontSize=12,
                        leading=20, spaceBefore=6, spaceAfter=4)
    title_style = ParagraphStyle("ArTitle", parent=body, fontName=FONT_B,
                                 fontSize=26, leading=36, alignment=TA_CENTER)
    center = ParagraphStyle("ArCenter", parent=body, alignment=TA_CENTER)

    class _Doc(BaseDocTemplate):
        """يسجّل العناوين في الفهرس بعد رسم كل عنصر."""

        def afterFlowable(self, flowable):
            level = getattr(flowable, "_toc_level", None)
            if level is None:
                return
            self.notify("TOCEntry", (level, getattr(flowable, "_toc_text", ""), self.page))

    bio = BytesIO()
    margin = 2.2 * cm
    doc = _Doc(
        bio, pagesize=A4,
        leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin,
        title=f"العرض الفني — {company_name}" if company_name else "العرض الفني",
    )
    avail = doc.width

    def on_page(canvas, doc_):
        if not include_page_numbers:
            return
        canvas.saveState()
        canvas.setFont(FONT, 9)
        canvas.drawCentredString(A4[0] / 2, margin / 2, _shape(str(doc_.page)))
        canvas.restoreState()

    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="body")
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=on_page)])

    def P(text, style, width=None):
        return Paragraph(_wrap_shaped(text, style.fontName, style.fontSize,
                                      width or avail), style)

    def H(text, style, level):
        """عنوان يُسجَّل في فهرس المحتويات."""
        p = P(text, style)
        p._toc_level = level
        p._toc_text = _shape(text)
        return p

    story = []

    # ── صفحة العنوان ──────────────────────────────────────────────────────────
    story.append(Spacer(1, 6 * cm))
    story.append(P("العرض الفني", title_style))
    story.append(Spacer(1, 1 * cm))
    if company_name:
        story.append(P(f"مقدَّم من: {company_name}", center))
    story.append(P(f"التاريخ: {datetime.date.today().strftime('%Y/%m/%d')}", center))
    story.append(PageBreak())

    # ── الفهرس ────────────────────────────────────────────────────────────────
    if include_toc:
        story.append(P("فهرس المحتويات", h1))
        toc = _rtl_toc_class()()
        toc.levelStyles = [
            ParagraphStyle("toc1", parent=body, fontName=FONT_B, fontSize=12,
                           leading=22, alignment=TA_RIGHT, spaceBefore=4),
            ParagraphStyle("toc2", parent=body, fontName=FONT, fontSize=10.5,
                           leading=18, alignment=TA_RIGHT, rightIndent=20),
        ]
        story.append(toc)
        story.append(PageBreak())

    def table_flowable(header: list, rows: list):
        header = list(reversed(header))
        rows = [list(reversed(r)) for r in rows]
        col_w = avail / max(len(header), 1)
        cell = ParagraphStyle("cell", parent=body, fontSize=9, leading=14, spaceAfter=0)
        cell_b = ParagraphStyle("cellb", parent=cell, fontName=FONT_B)
        data = [[P(str(c), cell_b, col_w - 10) for c in header]]
        for row in rows:
            data.append([P("" if v is None else str(v), cell, col_w - 10)
                         for v in row[: len(header)]])
        t = Table(data, colWidths=[col_w] * len(header), repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1E293B")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#94A3B8")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [colors.white, colors.HexColor("#F1F5F9")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
            ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        return t

    heading_styles = {1: h1, 2: h2, 3: h3}

    for idx, sec in enumerate(sections):
        kind = sec.get("kind")
        story.append(H(sec.get("title", ""), h1, level=0))

        if kind == "table_compliance":
            block = _df_to_table_block(df_compliance)
            story.append(table_flowable(*block) if block
                         else P("[لا توجد بيانات في جدول الامتثال]", body))
        elif kind == "table_boq":
            block = _df_to_table_block(df_boq)
            story.append(table_flowable(*block) if block
                         else P("[لا توجد بيانات في جدول الكميات]", body))
        else:
            content = (sec.get("content") or "").strip()
            if not content:
                story.append(P("[هذا القسم فارغ]", body))
            for block in parse_blocks(content):
                btype = block["type"]
                if btype == "heading":
                    style = heading_styles.get(min(block["level"], 3), h3)
                    story.append(H(block["text"], style, level=1))
                elif btype == "paragraph":
                    story.append(P(block["text"], body))
                elif btype in ("bullets", "numbered"):
                    for n, item in enumerate(block["items"], start=1):
                        marker = "•" if btype == "bullets" else f"{n}."
                        story.append(P(f"{item} {marker}", body))
                elif btype == "table":
                    story.append(table_flowable(block["header"], block["rows"]))
                    story.append(Spacer(1, 0.3 * cm))

        if idx < len(sections) - 1:
            story.append(PageBreak())

    # multiBuild يُشغّل بناءين حتى تُحلّ أرقام صفحات الفهرس
    doc.multiBuild(story)
    bio.seek(0)
    return bio

