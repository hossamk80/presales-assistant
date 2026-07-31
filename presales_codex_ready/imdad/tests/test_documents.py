"""اختبارات بناء مستندات Word و PDF بالاتجاهين."""
import pytest

SECTIONS_AR = [
    {"key": "cover", "title": "خطاب التقديم", "kind": "cover",
     "content": "نفيدكم برغبتنا في تقديم هذا العرض الفني."},
    {"key": "docinfo", "title": "إشعار السرية", "kind": "docinfo",
     "content": "هذا المستند سري."},
    {"key": "m", "title": "المنهجية الفنية", "kind": "ai", "content": (
        "## المنهجية\n\nنعتمد منهجية رشيقة.\n\n- بند أول\n- بند ثانٍ\n\n"
        "| المعيار | الوزن |\n|---|---|\n| الخبرة | 40 |\n"
    )},
    {"key": "t", "title": "جدول الامتثال", "kind": "table_compliance", "content": ""},
]

SECTIONS_EN = [
    {"key": "cover", "title": "Cover Letter", "kind": "cover",
     "content": "We are pleased to submit this technical proposal."},
    {"key": "m", "title": "Methodology", "kind": "ai",
     "content": "## Approach\n\nWe follow an agile method.\n\n- First\n- Second\n"},
    {"key": "t", "title": "Compliance Matrix", "kind": "table_compliance", "content": ""},
]


@pytest.fixture()
def fh():
    from utils import file_handler
    return file_handler


@pytest.fixture()
def df():
    import pandas as pd
    return pd.DataFrame({
        "المتطلب التقني": ["شهادة ISO 27001"],
        "الالتزام": ["نعم"],
    })


@pytest.mark.parametrize("rtl,sections", [(True, SECTIONS_AR), (False, SECTIONS_EN)])
def test_word_builds_valid_document(fh, df, rtl, sections, tmp_path):
    import docx

    bio = fh.build_word_document("شركة الإمداد", sections, df_compliance=df, rtl=rtl)
    out = tmp_path / "p.docx"
    out.write_bytes(bio.getvalue())

    doc = docx.Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    for sec in sections:
        assert sec["title"] in text
    assert len(doc.tables) >= 1


def test_table_column_order_follows_direction(fh, df):
    """
    في العربية يُقلب ترتيب الأعمدة ليبدأ العمود الأول من اليمين؛ وفي
    الإنجليزية يبقى كما هو. خلط الاتجاهين يُخرج جدولاً مقلوباً في المستند.
    """
    import docx

    sections = [{"key": "t", "title": "جدول", "kind": "table_compliance", "content": ""}]
    columns = list(df.columns)

    rtl_doc = docx.Document(fh.build_word_document("ش", sections, df_compliance=df, rtl=True))
    ltr_doc = docx.Document(fh.build_word_document("ش", sections, df_compliance=df, rtl=False))

    rtl_header = [c.text for c in rtl_doc.tables[0].rows[0].cells]
    ltr_header = [c.text for c in ltr_doc.tables[0].rows[0].cells]

    assert ltr_header == columns
    assert rtl_header == list(reversed(columns))


def test_word_without_toc_or_page_numbers(fh):
    import docx

    bio = fh.build_word_document(
        "شركة", SECTIONS_AR, include_toc=False, include_page_numbers=False
    )
    doc = docx.Document(bio)
    assert "فهرس المحتويات" not in "\n".join(p.text for p in doc.paragraphs)


def test_word_marks_empty_sections(fh):
    import docx

    sections = [{"key": "x", "title": "قسم", "kind": "ai", "content": ""}]
    doc = docx.Document(fh.build_word_document("ش", sections))
    assert "[هذا القسم فارغ]" in "\n".join(p.text for p in doc.paragraphs)


@pytest.mark.parametrize("rtl,sections", [(True, SECTIONS_AR), (False, SECTIONS_EN)])
def test_pdf_builds_valid_document(fh, df, rtl, sections):
    pytest.importorskip("reportlab")
    pytest.importorskip("arabic_reshaper")

    from pypdf import PdfReader

    bio = fh.build_pdf_document("شركة الإمداد", sections, df_compliance=df, rtl=rtl)
    data = bio.getvalue()
    assert data.startswith(b"%PDF")

    reader = PdfReader(bio)
    # صفحة عنوان + فهرس + قسم لكل عنصر على الأقل
    assert len(reader.pages) >= len(sections) + 2


def test_pdf_requires_arabic_font(fh, monkeypatch):
    """
    غياب خط عربي يجب أن يرفع خطأ صريحاً لا أن يُخرج ملفاً مشوّهاً بصمت.
    """
    pytest.importorskip("reportlab")
    monkeypatch.setattr(fh, "_find_arabic_font", lambda: (None, None))
    with pytest.raises(ImportError):
        fh.build_pdf_document("شركة", SECTIONS_AR)


def test_shape_is_identity_for_ltr(fh):
    pytest.importorskip("arabic_reshaper")
    assert fh._shape("Hello World", rtl=False) == "Hello World"


def test_shape_transforms_arabic_for_rtl(fh):
    pytest.importorskip("arabic_reshaper")
    shaped = fh._shape("مرحبا", rtl=True)
    assert shaped != "مرحبا"
    assert len(shaped) > 0


def test_df_to_table_block_handles_empty(fh):
    import pandas as pd

    assert fh._df_to_table_block(None) is None
    assert fh._df_to_table_block(pd.DataFrame()) is None
    header, rows = fh._df_to_table_block(pd.DataFrame({"أ": [1], "ب": [2]}))
    assert header == ["أ", "ب"]
    assert rows == [["1", "2"]]
