"""اختبارات بناء مستندات Word و PDF بالاتجاهين."""
import pytest

SECTIONS_AR = [
    {"key": "cover", "title": "خطاب التقديم", "kind": "cover",
     "content": "نفيدكم برغبتنا في تقديم هذا العرض الفني."},
    {"key": "docinfo", "title": "إشعار السرية", "kind": "docinfo",
     "content": "هذا المستند سري."},
    {"key": "m", "title": "المنهجية الفنية", "kind": "ai", "content": (
        "## المنهجية\n\nنعتمد منهجية رشيقة.\n\n- بند أول\n- بند ثانٍ\n\n"
        "| المعيار | الوزن |\n|---|---|\n| الخبرة | 40 |\n"
    )},
    {"key": "t", "title": "جدول الامتثال", "kind": "table_compliance", "content": ""},
]

SECTIONS_EN = [
    {"key": "cover", "title": "Cover Letter", "kind": "cover",
     "content": "We are pleased to submit this technical proposal."},
    {"key": "m", "title": "Methodology", "kind": "ai",
     "content": "## Approach\n\nWe follow an agile method.\n\n- First\n- Second\n"},
    {"key": "t", "title": "Compliance Matrix", "kind": "table_compliance", "content": ""},
]


@pytest.fixture()
def fh():
    from utils import file_handler
    return file_handler


@pytest.fixture()
def df():
    import pandas as pd
    return pd.DataFrame({
        "المتطلب التقني": ["شهادة ISO 27001"],
        "الالتزام": ["نعم"],
    })


@pytest.mark.parametrize("rtl,sections", [(True, SECTIONS_AR), (False, SECTIONS_EN)])
def test_word_builds_valid_document(fh, df, rtl, sections, tmp_path):
    import docx

    bio = fh.build_word_document("شركة الحلول المتقدمة", sections, df_compliance=df, rtl=rtl)
    out = tmp_path / "p.docx"
    out.write_bytes(bio.getvalue())

    doc = docx.Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    for sec in sections:
        assert sec["title"] in text
    assert len(doc.tables) >= 1


def test_table_column_order_follows_direction(fh, df):
    """
    في العربية يُقلب ترتيب الأعمدة ليبدأ العمود الأول من اليمين؛ وفي
    الإنجليزية يبقى كما هو. خلط الاتجاهين يُخرج جدولاً مقلوباً في المستند.
    """
    import docx

    sections = [{"key": "t", "title": "جدول", "kind": "table_compliance", "content": ""}]
    columns = list(df.columns)

    rtl_doc = docx.Document(fh.build_word_document("ش", sections, df_compliance=df, rtl=True))
    ltr_doc = docx.Document(fh.build_word_document("ش", sections, df_compliance=df, rtl=False))

    rtl_header = [c.text for c in rtl_doc.tables[0].rows[0].cells]
    ltr_header = [c.text for c in ltr_doc.tables[0].rows[0].cells]

    assert ltr_header == columns
    assert rtl_header == list(reversed(columns))


def test_word_without_toc_or_page_numbers(fh):
    import docx

    bio = fh.build_word_document(
        "شركة", SECTIONS_AR, include_toc=False, include_page_numbers=False
    )
    doc = docx.Document(bio)
    assert "فهرس المحتويات" not in "\n".join(p.text for p in doc.paragraphs)


def test_word_marks_empty_sections(fh):
    import docx

    sections = [{"key": "x", "title": "قسم", "kind": "ai", "content": ""}]
    doc = docx.Document(fh.build_word_document("ش", sections))
    assert "[هذا القسم فارغ]" in "\n".join(p.text for p in doc.paragraphs)


@pytest.mark.parametrize("rtl,sections", [(True, SECTIONS_AR), (False, SECTIONS_EN)])
def test_pdf_builds_valid_document(fh, df, rtl, sections):
    pytest.importorskip("reportlab")
    pytest.importorskip("arabic_reshaper")

    from pypdf import PdfReader

    bio = fh.build_pdf_document("شركة الحلول المتقدمة", sections, df_compliance=df, rtl=rtl)
    data = bio.getvalue()
    assert data.startswith(b"%PDF")

    reader = PdfReader(bio)
    # صفحة عنوان + فهرس + قسم لكل عنصر على الأقل
    assert len(reader.pages) >= len(sections) + 2


def test_pdf_requires_arabic_font(fh, monkeypatch):
    """
    غياب خط عربي يجب أن يرفع خطأ صريحاً لا أن يُخرج ملفاً مشوّهاً بصمت.
    """
    pytest.importorskip("reportlab")
    monkeypatch.setattr(fh, "_find_arabic_font", lambda: (None, None))
    with pytest.raises(ImportError):
        fh.build_pdf_document("شركة", SECTIONS_AR)


def test_shape_is_identity_for_ltr(fh):
    pytest.importorskip("arabic_reshaper")
    assert fh._shape("Hello World", rtl=False) == "Hello World"


def test_shape_transforms_arabic_for_rtl(fh):
    pytest.importorskip("arabic_reshaper")
    shaped = fh._shape("مرحبا", rtl=True)
    assert shaped != "مرحبا"
    assert len(shaped) > 0


def test_df_to_table_block_handles_empty(fh):
    import pandas as pd

    assert fh._df_to_table_block(None) is None
    assert fh._df_to_table_block(pd.DataFrame()) is None
    header, rows = fh._df_to_table_block(pd.DataFrame({"أ": [1], "ب": [2]}))
    assert header == ["أ", "ب"]
    assert rows == [["1", "2"]]


# ─── تاريخ الغلاف ونماذج الجهات (ب-4) ─────────────────────────────────────────
#
# المراسلة الحكومية السعودية تحمل التقويمين، وبعض الجهات ترفض عرضاً بغير
# نموذجها — رفضاً شكلياً لا علاقة له بجودة المحتوى.


def test_the_cover_carries_both_calendars_each_marked(fake_streamlit):
    """
    **الوسم `هـ` و `م` ليس زينة**: تاريخ هجري بلا وسم يُقرأ ميلادياً — وهو
    المزلق نفسه الذي عولج في 14-9. وعلى غلاف يقرؤه مُقيّم، الالتباس بين 1448
    و 2026 ليس تفصيلاً.
    """
    import datetime

    from utils import submission

    line = submission.cover_date(datetime.date(2026, 8, 8), rtl=True)

    assert "1448" in line and "هـ" in line
    assert "2026/08/08" in line and "م" in line


def test_the_english_cover_marks_both_too(fake_streamlit):
    import datetime

    from utils import submission

    line = submission.cover_date(datetime.date(2026, 8, 8), rtl=False)

    assert "AD" in line and "AH" in line


def test_a_missing_calendar_library_shows_gregorian_alone(fake_streamlit,
                                                          monkeypatch):
    """
    **إمّا تحويل صحيح أو لا تاريخ.** تاريخ هجري خاطئ على غلاف عرض حكومي أسوأ
    من غيابه: الغائب يُستدرَك، والخاطئ يُقرأ صحيحاً ويُبنى عليه. فبلا المكتبة
    يظهر الميلادي وحده — لا تقدير حسابي.
    """
    import builtins
    import datetime

    from utils import submission

    real_import = builtins.__import__

    def no_calendar(name, *args, **kwargs):
        if name in ("hijridate", "hijri_converter"):
            raise ImportError(name)
        return real_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", no_calendar)

    assert submission.gregorian_to_hijri(datetime.date(2026, 8, 8)) is None
    assert submission.cover_date(datetime.date(2026, 8, 8)) == "2026/08/08"


def test_the_generated_cover_actually_shows_the_hijri_date(fake_streamlit):
    """التحقّق من المستند المبنيّ فعلاً لا من الدالّة وحدها."""
    import io

    from docx import Document

    from utils.file_handler import build_word_document

    bio = build_word_document(
        company_name="شركة",
        sections=[{"kind": "ai", "title": "المنهجية", "content": "نص"}],
        entity_name="وزارة الصحة", rtl=True,
    )
    lines = [p.text for p in Document(io.BytesIO(bio.getvalue())).paragraphs[:8]]
    date_line = next(x for x in lines if x.startswith("التاريخ"))

    assert "هـ" in date_line


def test_an_entity_template_is_found_whatever_the_spelling(temp_db):
    """
    «وزارة الصحة» و«وزاره الصحه» جهة واحدة — بالتوحيد نفسه في ذاكرة العطاءات
    (12-7) لا بتوحيدٍ ثانٍ: قاعدتان تعنيان جهةً تُطابَق هنا ولا تُطابَق هناك.
    """
    temp_db.save_entity_template("وزارة الصحة", b"PK-form", filename="moh.docx")

    for spelling in ("وزارة الصحة", "وزاره الصحه ", "وزارة الصحه"):
        found = temp_db.entity_template(spelling)
        assert found is not None, spelling
        assert found["entity_label"] == "وزارة الصحة"

    assert temp_db.entity_template("وزارة النقل") is None


def test_saving_again_replaces_rather_than_duplicates(temp_db):
    """نموذجان لجهة واحدة يجعلان التصدير يختار أحدهما بلا قاعدة."""
    temp_db.save_entity_template("وزارة الصحة", b"v1", filename="a.docx")
    temp_db.save_entity_template("وزاره الصحه", b"v2", filename="b.docx")

    assert len(temp_db.list_entity_templates()) == 1
    assert temp_db.entity_template("وزارة الصحة")["template"] == b"v2"


def test_an_empty_entity_or_file_is_refused(temp_db):
    assert temp_db.save_entity_template("", b"x") is None
    assert temp_db.save_entity_template("جهة", b"") is None
    assert temp_db.list_entity_templates() == []


def test_the_listing_does_not_carry_the_file_payloads(temp_db):
    """القائمة تُعرض في كل رسم — تحميل الملفات لها يُثقلها بلا داعٍ."""
    temp_db.save_entity_template("جهة", b"PK" * 5000, filename="x.docx")
    row = temp_db.list_entity_templates()[0]

    assert "template" not in row
    assert row["size"] == 10000


def test_the_entity_form_wins_over_the_company_template(temp_db, fake_streamlit):
    """
    **شرط قبول ب-4**: التصدير يتبع نموذج الجهة حين يوجد.
    """
    from views import doc_builder

    temp_db.save_entity_template("وزارة الصحة", b"ENTITY-FORM")
    doc_builder.st.session_state["_project_entity"] = "وزارة الصحة"
    doc_builder.st.session_state["c_word_template_bytes"] = b"COMPANY"

    assert doc_builder._export_template() == b"ENTITY-FORM"


def test_the_company_template_covers_entities_with_no_form(temp_db, fake_streamlit):
    from views import doc_builder

    doc_builder.st.session_state["_project_entity"] = "جهة بلا نموذج"
    doc_builder.st.session_state["c_word_template_bytes"] = b"COMPANY"

    assert doc_builder._export_template() == b"COMPANY"


def test_the_user_can_force_the_company_template(temp_db, fake_streamlit):
    """نموذج جهة قديم أسوأ من غيابه — القرار يبقى بيد المستخدم."""
    from views import doc_builder

    temp_db.save_entity_template("وزارة الصحة", b"ENTITY-FORM")
    doc_builder.st.session_state["_project_entity"] = "وزارة الصحة"
    doc_builder.st.session_state["c_word_template_bytes"] = b"COMPANY"
    doc_builder.st.session_state["exp_force_company"] = True

    assert doc_builder._export_template() == b"COMPANY"


def test_the_entity_falls_back_to_what_the_tender_says(temp_db, fake_streamlit):
    """الجهة قد لا تُكتب يدوياً — تُقرأ ممّا استُخرج من الكرّاس."""
    from views import doc_builder

    temp_db.save_entity_template("وزارة النقل", b"NAQL-FORM")
    doc_builder.st.session_state["project_context"] = {"issuing_entity": "وزارة النقل"}

    assert doc_builder._entity_name() == "وزارة النقل"
    assert doc_builder._export_template() == b"NAQL-FORM"


# ─── الأشكال: مخططات العرض وصوره (ب-5) ────────────────────────────────────────
#
# العرض الفني بلا مخطط معماري يخسر درجات في «وضوح الحل»: لجنة الفحص تقرأ خمسين
# صفحة نصّاً لتفهم بنيةً يوضّحها شكل واحد.
#
# **والصورة يرفعها المستخدم ولا يولّدها النموذج**: النموذج لا يعرف معمارية الحلّ
# الفعلية، ومخططٌ مولَّد ادّعاءٌ عن الحلّ لا توضيحٌ له.


def _png(width: int = 60, height: int = 40) -> bytes:
    """صورة PNG صالحة صغيرة — نبنيها ولا نُرفق ملفاً ثنائياً بالاختبارات."""
    import struct
    import zlib

    def chunk(tag, data):
        payload = tag + data
        return struct.pack(">I", len(data)) + payload + struct.pack(
            ">I", zlib.crc32(payload))

    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    raw = b"".join(b"\x00" + bytes([80, 120, 160]) * width for _ in range(height))
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header)
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


_FIG_SECTIONS = [
    {"key": "intro", "title": "المقدمة", "kind": "ai", "content": "نص المقدمة."},
    {"key": "arch", "title": "المعمارية", "kind": "ai", "content": "انظر شكل 2."},
]


def test_a_file_is_judged_by_its_signature_not_its_name(fake_streamlit):
    """
    ملفٌّ سُمّي `.png` وهو مستند Word يمرّ من فحص الامتداد ثم يُسقط بناء
    المستند — بعيداً عن سببه بأيام. الفحص بالتوقيع.
    """
    from utils import figures

    assert figures.detect_image(_png()) == "image/png"
    assert figures.detect_image(b"PK\x03\x04 this is really a docx") is None
    assert figures.problem(b"PK\x03\x04 docx") == "fig.err_not_image"


def test_a_bad_upload_is_refused_at_upload_time(fake_streamlit):
    """صورة تالفة تُقبل صامتةً ثم تُسقط بناء المستند كلّه في يوم التسليم."""
    from utils import figures

    assert figures.problem(b"") == "fig.err_empty"
    assert figures.problem(b"\x89PNG\r\n\x1a\n" + b"x" * (figures.MAX_BYTES + 1)) \
        == "fig.err_too_big"
    assert figures.problem(_png()) is None


def test_figures_are_numbered_in_document_order(fake_streamlit):
    """
    القارئ يرى «شكل ٢» بعد «شكل ١» — والترتيب ترتيب المستند لا ترتيب الرفع.
    """
    from utils import figures

    uploaded_late_but_earlier_section = [
        {"id": 3, "section_key": "arch", "ordinal": 0, "caption": "معمارية"},
        {"id": 1, "section_key": "intro", "ordinal": 0, "caption": "نطاق"},
    ]
    out = figures.numbered(uploaded_late_but_earlier_section, _FIG_SECTIONS)

    assert [f["caption"] for f in out] == ["نطاق", "معمارية"]
    assert [f["label"] for f in out] == ["شكل 1", "شكل 2"]


def test_deleting_a_figure_leaves_no_gap(fake_streamlit):
    """
    الرقم يُحسب وقت البناء ولا يُخزَّن — نفس مبدأ ترقيم الملاحق (12-9).
    """
    from utils import figures

    remaining = [{"id": 2, "section_key": "arch", "ordinal": 0, "caption": "ب"}]
    out = figures.numbered(remaining, _FIG_SECTIONS)

    assert [f["label"] for f in out] == ["شكل 1"]


def test_a_figure_in_an_excluded_section_is_dropped(fake_streamlit):
    """
    القسم لا يخرج في المستند فشكله لا مكان له — وإدراجه يُنتج «شكل ٣» بلا شكل
    ٣ في النصّ.
    """
    from utils import figures

    sections = _FIG_SECTIONS + [{"key": "out", "title": "مستبعد", "include": False}]
    figs = [{"id": 1, "section_key": "out", "ordinal": 0, "caption": "مستبعد"}]

    assert figures.numbered(figs, sections) == []


def test_the_english_label_follows_the_output_language(fake_streamlit):
    from utils import figures

    out = figures.numbered(
        [{"id": 1, "section_key": "intro", "ordinal": 0}], _FIG_SECTIONS, rtl=False)

    assert out[0]["label"] == "Figure 1"


def test_word_and_pdf_carry_the_same_figures(fake_streamlit):
    """
    **شرط قبول ب-5**: الأشكال تخرج في المستند — والصيغتان لا تختلفان أمام لجنة
    الفتح.
    """
    import io

    from docx import Document

    from utils.file_handler import build_pdf_document, build_word_document

    image = _png()
    figs = [
        {"id": 1, "section_key": "intro", "ordinal": 0, "caption": "نطاق المشروع",
         "image": image},
        {"id": 2, "section_key": "arch", "ordinal": 0, "caption": "معمارية الحل",
         "image": image},
    ]

    bio = build_word_document(company_name="شركة", sections=_FIG_SECTIONS,
                              figures=figs, rtl=True)
    doc = Document(io.BytesIO(bio.getvalue()))
    captions = [p.text for p in doc.paragraphs if p.text.startswith("شكل")]

    assert captions == ["شكل 1: نطاق المشروع", "شكل 2: معمارية الحل"]
    assert len(doc.inline_shapes) == 2

    pdf = build_pdf_document(company_name="شركة", sections=_FIG_SECTIONS,
                             figures=figs, rtl=True)
    assert pdf.getvalue().startswith(b"%PDF-")


def test_a_corrupt_image_does_not_break_the_whole_document(fake_streamlit):
    """
    صورة أفلتت فحص الرفع تُتخطّى ولا تُسقط المستند كلّه في يوم التسليم —
    والنقص مرئي، بخلاف مستند لم يُبنَ أصلاً.
    """
    from utils.file_handler import build_word_document

    broken = [{"id": 1, "section_key": "intro", "ordinal": 0, "caption": "تالفة",
               "image": b"\x89PNG\r\n\x1a\nGARBAGE"}]

    bio = build_word_document(company_name="شركة", sections=_FIG_SECTIONS,
                              figures=broken, rtl=True)

    assert len(bio.getvalue()) > 0


def test_a_document_with_no_figures_is_unchanged(fake_streamlit):
    """البند إضافة لا تغيير: عرضٌ بلا أشكال يُبنى كما كان."""
    from utils.file_handler import build_word_document

    assert len(build_word_document(company_name="ش", sections=_FIG_SECTIONS,
                                   rtl=True).getvalue()) > 0


def test_a_caption_is_wanted_but_not_forced(fake_streamlit):
    """
    شكلٌ بلا تسمية يُجبر القارئ على استنتاج ما يراه، ولجنة الفحص لا تستنتج.
    يُنبَّه ولا يُمنع — قد يكون الشكل مفهوماً بذاته، والقرار للمستخدم.
    """
    from utils import figures

    out = figures.numbered(
        [{"id": 1, "section_key": "intro", "ordinal": 0, "caption": ""}],
        _FIG_SECTIONS)

    assert figures.missing_captions(out)
    assert figures.caption_text(out[0]) == "شكل 1"      # يخرج بالرقم وحده


def test_a_figure_the_text_never_mentions_is_flagged(fake_streamlit):
    """شكلٌ لا يشير إليه النصّ يبدو حشواً — تنبيهٌ لا شرط."""
    from utils import figures

    out = figures.numbered(
        [{"id": 1, "section_key": "intro", "ordinal": 0},
         {"id": 2, "section_key": "arch", "ordinal": 0}], _FIG_SECTIONS)

    assert figures.referenced_in("انظر شكل 2 للتفاصيل", out[1]) is True
    assert figures.referenced_in("نص المقدمة.", out[0]) is False


def test_the_listing_does_not_load_the_images(temp_db):
    """القائمة تُرسم في كل دورة — قراءة ميغابايت لها في كل مرة عبث."""
    pid = temp_db.create_project("م", {})
    temp_db.add_figure(pid, "intro", _png(), caption="نطاق")

    light = temp_db.list_figures(pid)[0]
    heavy = temp_db.list_figures(pid, with_images=True)[0]

    assert "image" not in light and light["size"] > 0
    assert heavy["image"].startswith(b"\x89PNG")


def test_figures_keep_their_order_within_a_section(temp_db):
    pid = temp_db.create_project("م", {})
    first = temp_db.add_figure(pid, "intro", _png(), caption="أ")
    second = temp_db.add_figure(pid, "intro", _png(), caption="ب")

    rows = temp_db.list_figures(pid)
    assert [r["id"] for r in rows] == [first, second]
    assert [r["ordinal"] for r in rows] == [0, 1]


def test_an_empty_image_is_refused_by_the_store(temp_db):
    pid = temp_db.create_project("م", {})

    assert temp_db.add_figure(pid, "intro", b"") is None
    assert temp_db.list_figures(pid) == []


def test_deleting_a_tender_takes_its_figures(temp_db):
    pid = temp_db.create_project("م", {})
    temp_db.add_figure(pid, "intro", _png())
    temp_db.add_figure(pid, "arch", _png())

    assert temp_db.delete_project_figures(pid) == 2
    assert temp_db.list_figures(pid) == []
